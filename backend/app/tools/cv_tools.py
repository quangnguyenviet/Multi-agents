import io
import json
import unicodedata
import uuid

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from langchain_core.tools import tool

from app.services.cv_service import extract_cv_data
from app.repositories.blob_store import make_blob_store
from app.tools.file_tools import _upload_store

_cv_json_store = make_blob_store("cv_json")
_cv_docx_store = make_blob_store("cv_docx")


@tool
def read_cv_file(file_id: str) -> str:
    """Trích xuất thông tin CV có cấu trúc từ file PDF đã upload, trả về JSON gồm:
    personal_info, summary, experience, education, skills, languages, certifications.
    Chỉ dùng khi đã xác định file là CV/hồ sơ xin việc (sau khi đọc bằng read_file_content).
    Sau khi đọc xong, hãy nhớ file_id để dùng lại với generate_cv_word_file.
    Người dùng có thể yêu cầu đổi mẫu hoặc ngôn ngữ nhiều lần — chỉ cần truyền lại file_id."""
    cached = _cv_json_store.get(file_id)
    if cached:
        return cached

    pdf_bytes = _upload_store.get(file_id)
    if not pdf_bytes:
        return f"Lỗi: Không tìm thấy file với ID '{file_id}'. Vui lòng upload lại file CV."
    cv_data = extract_cv_data(pdf_bytes)
    result = json.dumps(cv_data, ensure_ascii=False, indent=2)
    _cv_json_store[file_id] = result
    return result


# ── Color palette ─────────────────────────────────────────────────────────────
_C_NAVY      = RGBColor(0x0f, 0x34, 0x60)
_C_ACCENT    = RGBColor(0x5c, 0x6b, 0xc0)
_C_DARK      = RGBColor(0x1a, 0x1a, 0x2e)
_C_WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
_C_HEADER_BG = "2C3E50"


def _lv(para, label: str, value: str):
    """Label (bold navy) + value trong cùng paragraph."""
    lr = para.add_run(label)
    lr.bold = True; lr.font.size = Pt(10); lr.font.color.rgb = _C_NAVY
    vr = para.add_run(value or "")
    vr.font.size = Pt(10); vr.font.color.rgb = _C_DARK


def _para_bottom_border(para):
    """Thêm đường kẻ ngang mỏng phía dưới paragraph (dùng trong cell)."""
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), "4")
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), "AAAAAA")
    pBdr.append(bot)
    pPr.append(pBdr)


def _set_cell_bg(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _strip_accents(s: str) -> str:
    nfd = unicodedata.normalize("NFD", s or "")
    return "".join(c for c in nfd if unicodedata.category(c) != "Mn").lower()


def _date_range(a, b) -> str:
    a = (a or "").strip()
    b = (b or "").strip()
    if not a and not b:
        return ""
    return f"{a} - {b}".strip(" -")


_LABELS = {
    "vi": {
        "t1_title":       "BẢN LÝ LỊCH CHUYÊN MÔN CỦA NHÂN SỰ CHỦ CHỐT",
        "t1_position":    "Vị trí: ",
        "t1_personal":    "Thông tin nhân sự",
        "t1_name":        "Tên: ",
        "t1_phone":       "Điện thoại",
        "t1_dob":         "Ngày sinh: ",
        "t1_address":     "Địa chỉ",
        "t1_gender":      "Giới tính",
        "t1_education":   "Trình độ chuyên môn: ",
        "t1_experience":  "Kinh nghiệm chuyên môn",
        "t1_project":     "Tên Dự án: ",
        "t1_job_pos":     "Vị trí công việc: ",
        "t1_tasks":       "Công việc thực hiện:",
        "t1_tech_used":   "Công nghệ sử dụng: ",
        "t1_skills":      "Kỹ năng",
        "t1_technical":   "Kỹ thuật: ",
        "t1_soft":        "Kỹ năng mềm: ",
        "t1_languages":   "Ngoại ngữ",
        "t1_certs":       "Chứng chỉ",
        "t2_fullname":    "HỌ VÀ TÊN",
        "t2_position":    "VỊ TRÍ",
        "t2_overview":    "TỔNG QUAN",
        "t2_education":   "HỌC VẤN",
        "t2_languages":   "NGÔN NGỮ",
        "t2_lang_levels": ["Thành thạo", "Khá", "Trung bình"],
        "t2_tech":        "CÔNG NGHỆ",
        "t2_os":          "Hệ điều hành",
        "t2_core":        "Công nghệ chính",
        "t2_db":          "Cơ sở dữ liệu",
        "t2_tools":       "Công cụ",
        "t2_methods":     "Phương pháp",
        "t2_experience":  "KINH NGHIỆM LÀM VIỆC ({n} dự án)",
        "t2_project_n":   "Dự án {i}",
        "t2_period":      "Thời gian",
        "t2_pos":         "Vị trí",
        "t2_teamsize":    "Quy mô dự án",
        "t2_desc":        "Mô tả",
        "t2_tasks":       "Nhiệm vụ",
        "t2_tech_used":   "Công nghệ",
    },
    "en": {
        "t1_title":       "PROFESSIONAL RESUME OF KEY PERSONNEL",
        "t1_position":    "Position: ",
        "t1_personal":    "Personal Information",
        "t1_name":        "Name: ",
        "t1_phone":       "Phone",
        "t1_dob":         "Date of Birth: ",
        "t1_address":     "Address",
        "t1_gender":      "Gender",
        "t1_education":   "Education: ",
        "t1_experience":  "Professional Experience",
        "t1_project":     "Project: ",
        "t1_job_pos":     "Position: ",
        "t1_tasks":       "Responsibilities:",
        "t1_tech_used":   "Technologies: ",
        "t1_skills":      "Skills",
        "t1_technical":   "Technical: ",
        "t1_soft":        "Soft Skills: ",
        "t1_languages":   "Languages",
        "t1_certs":       "Certifications",
        "t2_fullname":    "FULL NAME",
        "t2_position":    "POSITION",
        "t2_overview":    "OVERVIEW",
        "t2_education":   "EDUCATION",
        "t2_languages":   "LANGUAGES",
        "t2_lang_levels": ["Fluent", "Good", "Basic"],
        "t2_tech":        "TECHNOLOGIES",
        "t2_os":          "Operating Systems",
        "t2_core":        "Core Technologies",
        "t2_db":          "Databases",
        "t2_tools":       "Tools",
        "t2_methods":     "Methodologies",
        "t2_experience":  "WORK EXPERIENCE ({n} projects)",
        "t2_project_n":   "Project {i}",
        "t2_period":      "Period",
        "t2_pos":         "Position",
        "t2_teamsize":    "Team Size",
        "t2_desc":        "Description",
        "t2_tasks":       "Responsibilities",
        "t2_tech_used":   "Technologies",
    },
}


def _build_docx_template1(cv_data: dict, labels: dict) -> Document:
    doc = Document()
    for sec in doc.sections:
        sec.top_margin    = Cm(1.5)
        sec.bottom_margin = Cm(1.5)
        sec.left_margin   = Cm(2.0)
        sec.right_margin  = Cm(2.0)

    pi         = cv_data.get("personal_info") or {}
    education  = cv_data.get("education") or []
    experience = cv_data.get("experience") or []
    skills     = cv_data.get("skills") or {}
    languages  = cv_data.get("languages") or []
    certs      = cv_data.get("certifications") or []

    W0, W1, W2 = Cm(3.8), Cm(9.0), Cm(4.2)

    tbl = doc.add_table(rows=0, cols=3)
    tbl.style = "Table Grid"
    tbl.allow_autofit = False

    def _set_w(row):
        row.cells[0].width = W0
        row.cells[1].width = W1
        row.cells[2].width = W2

    def _full_row(text="", bold=False, centered=False, size=10, color=None):
        row = tbl.add_row()
        _set_w(row)
        cell = row.cells[0].merge(row.cells[2])
        p = cell.paragraphs[0]
        if centered:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if text:
            r = p.add_run(text)
            r.bold = bold
            r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
        return cell

    _full_row(labels["t1_title"], bold=True, centered=True, size=13, color=_C_NAVY)

    pos = (experience[0].get("position") or "") if experience else ""
    _full_row(f"{labels['t1_position']}{pos}", bold=True, centered=True, size=11, color=_C_ACCENT)

    row = tbl.add_row()
    _set_w(row)
    c0, c1, c2 = row.cells

    r0 = c0.paragraphs[0].add_run(labels["t1_personal"])
    r0.bold = True; r0.font.size = Pt(10); r0.font.color.rgb = _C_NAVY
    c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    c1.paragraphs[0].add_run(f"{labels['t1_name']}{pi.get('full_name') or ''}").font.size = Pt(10)
    for key, label in [("email", "Email"), ("phone", labels["t1_phone"]), ("linkedin", "LinkedIn")]:
        if pi.get(key):
            c1.add_paragraph(f"{label}: {pi[key]}").runs[0].font.size = Pt(10)

    c2.paragraphs[0].add_run(f"{labels['t1_dob']}{pi.get('date_of_birth') or ''}").font.size = Pt(10)
    for key, label in [("address", labels["t1_address"]), ("gender", labels["t1_gender"])]:
        if pi.get(key):
            c2.add_paragraph(f"{label}: {pi[key]}").runs[0].font.size = Pt(10)

    if education:
        edu = education[0]
        cell = _full_row()
        p = cell.paragraphs[0]
        lr = p.add_run(labels["t1_education"])
        lr.bold = True; lr.font.size = Pt(10); lr.font.color.rgb = _C_NAVY
        parts = [x for x in [edu.get("degree"), edu.get("field")] if x]
        edu_text = " ".join(parts)
        if edu.get("institution"):
            edu_text += (" - " if edu_text else "") + edu["institution"]
        if edu.get("gpa"):
            edu_text += f" (GPA: {edu['gpa']})"
        p.add_run(edu_text).font.size = Pt(10)

    if experience:
        _full_row(labels["t1_experience"], bold=True, centered=False, size=11, color=_C_NAVY)

    for exp in experience:
        row = tbl.add_row()
        _set_w(row)
        dc = row.cells[0]
        ic = row.cells[1].merge(row.cells[2])

        dc.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        dp = dc.paragraphs[0]
        dp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        dr = dp.add_run(f"{exp.get('start_date', '')} - {exp.get('end_date', '')}")
        dr.font.size = Pt(9.5); dr.font.color.rgb = _C_DARK

        p1 = ic.paragraphs[0]
        _lv(p1, labels["t1_project"], exp.get("company") or "")
        _para_bottom_border(p1)

        p2 = ic.add_paragraph()
        _lv(p2, labels["t1_job_pos"], exp.get("position") or "")
        _para_bottom_border(p2)

        p3 = ic.add_paragraph()
        tr = p3.add_run(labels["t1_tasks"])
        tr.bold = True; tr.font.size = Pt(10); tr.font.color.rgb = _C_NAVY

        for bullet in (exp.get("description") or []):
            bp = ic.add_paragraph(f"•  {bullet}")
            bp.paragraph_format.left_indent = Inches(0.15)
            for br in bp.runs:
                br.font.size = Pt(9.5); br.font.color.rgb = _C_DARK

        if exp.get("technologies"):
            _lv(ic.add_paragraph(), labels["t1_tech_used"], exp["technologies"])

    tech = skills.get("technical") or []
    soft = skills.get("soft") or []
    if tech or soft:
        _full_row(labels["t1_skills"], bold=True, size=11, color=_C_NAVY)
        if tech:
            cell = _full_row()
            _lv(cell.paragraphs[0], labels["t1_technical"], "  ·  ".join(tech))
        if soft:
            cell = _full_row()
            _lv(cell.paragraphs[0], labels["t1_soft"], "  ·  ".join(soft))

    if languages:
        _full_row(labels["t1_languages"], bold=True, size=11, color=_C_NAVY)
        for lang in languages:
            cell = _full_row()
            cell.paragraphs[0].add_run(
                f"{lang.get('language') or ''}: {lang.get('level') or ''}"
            ).font.size = Pt(10)

    if certs:
        _full_row(labels["t1_certs"], bold=True, size=11, color=_C_NAVY)
        for cert in certs:
            parts = [x for x in [cert.get("name"), cert.get("issuer"), cert.get("date")] if x]
            cell = _full_row()
            cell.paragraphs[0].add_run("  |  ".join(parts)).font.size = Pt(10)

    return doc


_W2_LABEL, _W2_VALUE = Cm(4.5), Cm(12.5)
_W2_SECTION        = Cm(17.0)
_W2_LANG_NAME      = Cm(5.0)
_W2_LANG_LEVEL     = Cm(4.0)


def _t2_section_heading(doc: Document, text: str):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    cell = tbl.rows[0].cells[0]
    cell.width = _W2_SECTION
    _set_cell_bg(cell, _C_HEADER_BG)
    r = cell.paragraphs[0].add_run(text)
    r.bold = True; r.font.size = Pt(11); r.font.color.rgb = _C_WHITE
    return tbl


def _t2_kv(tbl, label: str, value: str, header: bool = False, bold_label: bool = True, size: int = 10):
    row = tbl.add_row()
    c0, c1 = row.cells
    c0.width = _W2_LABEL; c1.width = _W2_VALUE
    r = c0.paragraphs[0].add_run(label)
    r.bold = bold_label; r.font.size = Pt(size)
    if header:
        r.font.color.rgb = _C_WHITE
        _set_cell_bg(c0, _C_HEADER_BG)
    else:
        r.font.color.rgb = _C_NAVY
    vr = c1.paragraphs[0].add_run(value or "")
    vr.font.size = Pt(size); vr.font.color.rgb = _C_DARK
    return row


def _t2_kv_multiline(tbl, label: str, lines: list, header: bool = False):
    row = tbl.add_row()
    c0, c1 = row.cells
    c0.width = _W2_LABEL; c1.width = _W2_VALUE
    r = c0.paragraphs[0].add_run(label)
    r.bold = False; r.font.size = Pt(10)
    if header:
        r.font.color.rgb = _C_WHITE
        _set_cell_bg(c0, _C_HEADER_BG)
        c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    else:
        r.font.color.rgb = _C_NAVY
    first = True
    for ln in lines:
        p = c1.paragraphs[0] if first else c1.add_paragraph()
        run = p.add_run(ln)
        run.font.size = Pt(9.5); run.font.color.rgb = _C_DARK
        first = False
    return row


def _t2_language_row(tbl, lang: dict):
    name = lang.get("language") or ""
    level = _strip_accents(lang.get("level") or "")

    idx = None
    if any(k in level for k in ["thanh thao", "fluent", "native", "ban ngu", "tot", "cao", "proficient", "advanced"]):
        idx = 0
    elif any(k in level for k in ["kha", "good", "upper", "intermediate"]):
        idx = 1
    elif any(k in level for k in ["trung binh", "basic", "average", "co ban", "beginner", "elementary"]):
        idx = 2

    row = tbl.add_row()
    cells = row.cells
    cells[0].width = _W2_LANG_NAME
    r0 = cells[0].paragraphs[0].add_run(name)
    r0.font.size = Pt(10); r0.font.color.rgb = _C_DARK
    for i in range(3):
        cells[i + 1].width = _W2_LANG_LEVEL
        txt = "x" if i == idx else ""
        r = cells[i + 1].paragraphs[0].add_run(txt)
        r.font.size = Pt(10); r.font.color.rgb = _C_DARK
        cells[i + 1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if idx is None and lang.get("level"):
        cells[1].paragraphs[0].add_run(f"  [{lang['level']}]").font.size = Pt(9)
    return row


def _build_docx_template2(cv_data: dict, labels: dict) -> Document:
    doc = Document()
    for sec in doc.sections:
        sec.top_margin    = Cm(1.5)
        sec.bottom_margin = Cm(1.5)
        sec.left_margin   = Cm(2.0)
        sec.right_margin  = Cm(2.0)

    pi             = cv_data.get("personal_info") or {}
    summary        = cv_data.get("summary") or ""
    summary_points = cv_data.get("summary_points") or []
    experience     = cv_data.get("experience") or []
    education      = cv_data.get("education") or []
    skills         = cv_data.get("skills") or {}
    tech_stack     = skills.get("tech_stack") or {}
    languages      = cv_data.get("languages") or []

    position = (experience[0].get("position") or "") if experience else ""

    # 1. Personal info
    info = doc.add_table(rows=0, cols=2)
    info.style = "Table Grid"
    _t2_kv(info, labels["t2_fullname"], pi.get("full_name") or "", header=True, bold_label=True, size=11)
    _t2_kv(info, labels["t2_position"], position, header=True, bold_label=True, size=11)
    doc.add_paragraph()

    # 2. Summary / Overview
    text = summary.strip() or " ".join(s.strip() for s in summary_points if s.strip())
    if text:
        _t2_section_heading(doc, f"> {labels['t2_overview']}")
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.size = Pt(10); r.font.color.rgb = _C_DARK
        doc.add_paragraph()

    # 3. Education
    if education:
        _t2_section_heading(doc, f"> {labels['t2_education']}")
        edu_tbl = doc.add_table(rows=0, cols=2)
        edu_tbl.style = "Table Grid"
        for edu in education:
            period = _date_range(edu.get("start_date"), edu.get("end_date")) or "—"
            major  = " ".join(x for x in [edu.get("field"), edu.get("degree")] if x)
            institution = edu.get("institution") or ""
            erow = edu_tbl.add_row()
            c0, c1 = erow.cells
            c0.width = _W2_LABEL; c1.width = _W2_VALUE
            _set_cell_bg(c0, _C_HEADER_BG)
            rp = c0.paragraphs[0].add_run(period)
            rp.font.size = Pt(10); rp.font.color.rgb = _C_WHITE
            c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if major:
                rm = c1.paragraphs[0].add_run(major)
                rm.font.size = Pt(10); rm.font.color.rgb = _C_DARK
            if institution:
                inst_p = c1.add_paragraph() if major else c1.paragraphs[0]
                ri = inst_p.add_run(institution)
                ri.bold = True; ri.font.size = Pt(10); ri.font.color.rgb = _C_DARK
        doc.add_paragraph()

    # 4. Languages
    if languages:
        _t2_section_heading(doc, f"> {labels['t2_languages']}")
        lang_tbl = doc.add_table(rows=0, cols=4)
        lang_tbl.style = "Table Grid"
        hdr_row = lang_tbl.add_row()
        for j, htext in enumerate([labels["t2_languages"], *labels["t2_lang_levels"]]):
            hc = hdr_row.cells[j]
            hc.width = _W2_LANG_NAME if j == 0 else _W2_LANG_LEVEL
            _set_cell_bg(hc, _C_HEADER_BG)
            hr = hc.paragraphs[0].add_run(htext)
            hr.bold = True; hr.font.size = Pt(10); hr.font.color.rgb = _C_WHITE
            hc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for lang in languages:
            _t2_language_row(lang_tbl, lang)
        doc.add_paragraph()

    # 5. Technologies
    groups = [
        (labels["t2_os"],      tech_stack.get("operating_systems")),
        (labels["t2_core"],    tech_stack.get("core") or skills.get("technical")),
        (labels["t2_db"],      tech_stack.get("databases")),
        (labels["t2_tools"],   tech_stack.get("tools")),
        (labels["t2_methods"], tech_stack.get("methodologies")),
    ]
    if any(vals for _, vals in groups):
        _t2_section_heading(doc, f"> {labels['t2_tech']}")
        tech_tbl = doc.add_table(rows=0, cols=2)
        tech_tbl.style = "Table Grid"
        for lbl, vals in groups:
            if vals:
                trow = tech_tbl.add_row()
                c0, c1 = trow.cells
                c0.width = _W2_LABEL; c1.width = _W2_VALUE
                _set_cell_bg(c0, _C_HEADER_BG)
                lr = c0.paragraphs[0].add_run(lbl)
                lr.bold = False; lr.font.size = Pt(10); lr.font.color.rgb = _C_WHITE
                first = True
                for val in vals:
                    tp = c1.paragraphs[0] if first else c1.add_paragraph()
                    tp.add_run(f"- {val}").font.size = Pt(10)
                    for tr in tp.runs:
                        tr.font.color.rgb = _C_DARK
                    first = False
        doc.add_paragraph()

    # 6. Work experience
    if experience:
        _t2_section_heading(doc, f"> {labels['t2_experience'].format(n=len(experience))}")
        for i, exp in enumerate(experience, 1):
            etbl = doc.add_table(rows=0, cols=2)
            etbl.style = "Table Grid"
            _t2_kv(etbl, labels["t2_project_n"].format(i=i), exp.get("company") or "",
                   header=True, bold_label=True)
            period = _date_range(exp.get("start_date"), exp.get("end_date"))
            if period:
                _t2_kv(etbl, labels["t2_period"], period, header=True, bold_label=False)
            if exp.get("position"):
                _t2_kv(etbl, labels["t2_pos"], exp["position"], header=True, bold_label=False)
            if exp.get("team_size"):
                _t2_kv(etbl, labels["t2_teamsize"], str(exp["team_size"]), header=True, bold_label=False)
            if exp.get("overview"):
                _t2_kv(etbl, labels["t2_desc"], exp["overview"], header=True, bold_label=False)
            tasks = exp.get("description") or []
            if tasks:
                _t2_kv_multiline(etbl, labels["t2_tasks"], [f"- {t}" for t in tasks], header=True)
            if exp.get("technologies"):
                _t2_kv(etbl, labels["t2_tech_used"], exp["technologies"], header=True, bold_label=False)
            doc.add_paragraph()

    return doc


@tool
def generate_cv_word_file(file_id: str, template_id: str = "1", language: str = "vi") -> str:
    """Tạo file CV định dạng Word (.docx) từ file_id của CV đã đọc bằng read_cv_file.

    file_id: ID của file CV đã upload (từ thẻ [Attached file: '...', file_id=<id>]).
             read_cv_file phải được gọi trước với cùng file_id này.

    template_id:
      "1" = Bản Lý Lịch Chuyên Môn (bảng gọn, kinh nghiệm dạng thời gian | nội dung).
      "2" = Hồ sơ năng lực chi tiết (TỔNG QUAN, CÔNG NGHỆ phân nhóm, từng dự án).

    language:
      "vi" = Tiêu đề và đề mục bằng tiếng Việt (mặc định).
      "en" = Tiêu đề và đề mục bằng tiếng Anh.

    Dùng khi người dùng yêu cầu xuất CV ra Word/.docx. Có thể gọi lại nhiều lần với
    cùng file_id để đổi mẫu hoặc ngôn ngữ mà không cần upload lại file."""
    cached = _cv_json_store.get(file_id)
    if not cached:
        return f"Lỗi: Không tìm thấy dữ liệu CV cho file_id '{file_id}'. Hãy gọi read_cv_file trước."
    try:
        cv_data = json.loads(cached)
    except json.JSONDecodeError as e:
        return f"Lỗi JSON không hợp lệ: {e}"

    lang_key = "en" if str(language).strip().lower() == "en" else "vi"
    labels = _LABELS[lang_key]
    builder = _build_docx_template2 if str(template_id).strip() == "2" else _build_docx_template1
    doc = builder(cv_data, labels)

    buf = io.BytesIO()
    doc.save(buf)
    docx_id = uuid.uuid4().hex[:8]
    _cv_docx_store[docx_id] = buf.getvalue()
    return f"File Word đã tạo thành công. __artifact__:docx:{docx_id}"
